"""Final validation of the reviewed guideline and related HTML links."""
from pathlib import Path
from urllib.parse import unquote
from bs4 import BeautifulSoup
from docx import Document
import pymupdf
import json

ROOT = Path(__file__).resolve().parent
errors = []
for file in ROOT.glob("*.html"):
    soup = BeautifulSoup(file.read_text(encoding="utf-8"), "html.parser")
    for tag, attribute in [("a", "href"), ("link", "href"), ("script", "src"), ("img", "src")]:
        for node in soup.find_all(tag):
            value = node.get(attribute, "")
            if value and not value.startswith(("http", "#", "mailto:", "tel:", "data:")):
                target = (file.parent / unquote(value.split("#")[0])).resolve()
                if not target.exists():
                    errors.append(f"{file.name}: missing {value}")

source = ROOT / "SAMC_HD_Master_Assets_SOUTH_APOLLO_v2 (1)" / "SAMC_Brand_Guidelines_v2_2_South_Apollo_HD.docx"
reviewed = ROOT / "SAMC_Brand_Guidelines_v2_2_South_Apollo_HD_REVIEWED.docx"
doc = Document(reviewed)
assert source.stat().st_size == 1912796, "Original benchmark changed"
assert len(doc.inline_shapes) == 1, "Crest missing"
assert len(doc.tables) == 20, "Unexpected table count"
assert all("v2.1" not in p.text for p in doc.paragraphs), "Old version remains in body"
assert all("v2.1" not in p.text for s in doc.sections for p in s.footer.paragraphs), "Old footer version"
assert all(t.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for t in list(doc.tables)[:6] + list(doc.tables)[7:]), "Missing repeated table header"
assert doc.core_properties.title
pdf = pymupdf.open(ROOT / "guideline-audit" / "reviewed.pdf")
assert len(pdf) == 9
assert all(len(page.get_text().strip()) > 500 for page in pdf), "Empty page remains"
assert len(pdf[0].get_images()) == 1
assert not errors, errors
print(json.dumps({
    "original_preserved": True,
    "reviewed_docx": reviewed.name,
    "pages": len(pdf),
    "content_pages": len(pdf),
    "crest_embedded": True,
    "version": "v2.2",
    "html_local_references": "PASS",
    "tables": len(doc.tables),
    "production_status": "Reviewed copy; institutional approval still required"
}, indent=2))
