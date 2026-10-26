"""Read-only audit of the supplied SAMC DOCX brand guideline."""
from pathlib import Path
from collections import Counter
from docx import Document
from docx.shared import Emu
import json
import zipfile
import re

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "SAMC_HD_Master_Assets_SOUTH_APOLLO_v2 (1)" / "SAMC_Brand_Guidelines_v2_2_South_Apollo_HD.docx"
OUT = ROOT / "guideline-audit"
OUT.mkdir(exist_ok=True)

doc = Document(SOURCE)
lines = []
fonts = Counter()
font_sizes = Counter()
colors = Counter()
empty_paragraphs = 0

for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if not text:
        empty_paragraphs += 1
        continue
    lines.append(f"P{index} [{paragraph.style.name}] {text}")
    for run in paragraph.runs:
        if run.font.name:
            fonts[run.font.name] += 1
        if run.font.size:
            font_sizes[str(round(run.font.size.pt, 2))] += 1
        if run.font.color and run.font.color.rgb:
            colors[str(run.font.color.rgb)] += 1

for table_index, table in enumerate(doc.tables):
    lines.append(f"\nTABLE {table_index}")
    for row_index, row in enumerate(table.rows):
        lines.append(f"R{row_index}: " + " | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    if run.font.size:
                        font_sizes[str(round(run.font.size.pt, 2))] += 1
                    if run.font.color and run.font.color.rgb:
                        colors[str(run.font.color.rgb)] += 1

headers_footers = []
for index, section in enumerate(doc.sections):
    headers_footers.append({
        "section": index,
        "header": " | ".join(p.text for p in section.header.paragraphs),
        "footer": " | ".join(p.text for p in section.footer.paragraphs),
        "page_mm": [round(section.page_width.mm, 1), round(section.page_height.mm, 1)],
        "margins_mm": [round(section.left_margin.mm, 1), round(section.right_margin.mm, 1), round(section.top_margin.mm, 1), round(section.bottom_margin.mm, 1)],
    })

with zipfile.ZipFile(SOURCE) as package:
    names = package.namelist()
    media = [{"path": n, "bytes": package.getinfo(n).file_size} for n in names if n.startswith("word/media/")]
    rels = package.read("word/_rels/document.xml.rels").decode("utf-8")
    external_links = re.findall(r'Target="(https?[^\"]+)', rels)
    xml = package.read("word/document.xml").decode("utf-8")
    drawings = xml.count("<w:drawing>")

report = {
    "source": str(SOURCE),
    "bytes": SOURCE.stat().st_size,
    "core_properties": {
        key: str(getattr(doc.core_properties, key))
        for key in ["title", "subject", "author", "last_modified_by", "created", "modified", "revision"]
    },
    "paragraphs": len(doc.paragraphs),
    "empty_paragraphs": empty_paragraphs,
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "drawings": drawings,
    "sections": headers_footers,
    "explicit_fonts": fonts,
    "font_sizes_pt": font_sizes,
    "font_colors": colors,
    "external_links": external_links,
    "media": media,
    "embedded_objects": [n for n in names if "embeddings/" in n],
}
(OUT / "metadata.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=dict), encoding="utf-8")
(OUT / "text.txt").write_text("\n".join(lines), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2, default=dict))
print("\nExtracted:", OUT / "text.txt")
