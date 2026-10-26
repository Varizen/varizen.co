"""Create a corrected copy of the supplied guideline; never overwrite the benchmark."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "SAMC_HD_Master_Assets_SOUTH_APOLLO_v2 (1)" / "SAMC_Brand_Guidelines_v2_2_South_Apollo_HD.docx"
OUT = ROOT / "SAMC_Brand_Guidelines_v2_2_South_Apollo_HD_REVIEWED.docx"
CREST = ROOT / "SAMC_HD_Master_Assets_SOUTH_APOLLO_v2 (1)" / "HD_Assets" / "01_SAMC_Full_Color_Transparent_HD.png"
doc = Document(SOURCE)


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def set_east_asia(run, family):
    run.font.name = family
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + attr), family)


def set_size(paragraph, value):
    for run in paragraph.runs:
        run.font.size = Pt(value)
        set_east_asia(run, "Hind Siliguri" if any("\u0980" <= c <= "\u09ff" for c in run.text) else "Aptos")


# Version and document metadata.
for paragraph in all_paragraphs(doc):
    for run in paragraph.runs:
        if "v2.1" in run.text:
            run.text = run.text.replace("v2.1", "v2.2")
for section in doc.sections:
    for part in (section.footer, section.first_page_footer, section.even_page_footer):
        for paragraph in part.paragraphs:
            for run in paragraph.runs:
                if "v2.1" in run.text:
                    run.text = run.text.replace("v2.1", "v2.2")

props = doc.core_properties
props.title = "SAMC Brand Guidelines — Institutional Crest Edition v2.2"
props.subject = "South Apollo Medical College & Hospital: identity, bilingual, digital and governance standards"
props.author = "South Apollo Medical College & Hospital"
props.keywords = "SAMC, South Apollo, brand guidelines, institutional crest, English, Bangla"
props.comments = "Reviewed copy; original benchmark preserved. Factual, legal and clinical approval remain institutional responsibilities."
props.category = "Brand governance"
props.revision = 2
props.modified = datetime.now(timezone.utc)

# Use proper heading styles for navigation/accessibility.
heading_map = {
    1: "Title", 2: "Subtitle", 7: "Heading 1", 13: "Heading 1", 19: "Heading 1",
    29: "Heading 1", 35: "Heading 1", 41: "Heading 1", 55: "Heading 1", 61: "Heading 1",
    67: "Heading 1", 72: "Heading 1", 77: "Heading 1", 82: "Heading 1", 98: "Heading 1",
}
for i, style in heading_map.items():
    if i < len(doc.paragraphs):
        doc.paragraphs[i].style = doc.styles[style]
for i, paragraph in enumerate(doc.paragraphs):
    if paragraph.text.startswith(tuple(f"{n:02d} /" for n in range(12))) or paragraph.text == "APPENDIX":
        paragraph.style = doc.styles["Heading 2"]

# Avoid splitting headings from their following content.
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading") or paragraph.text.startswith(tuple(f"{n:02d} /" for n in range(12))):
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.keep_together = True

# Table accessibility/pagination: repeat headers, don't split individual rows.
for table in doc.tables:
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_size(paragraph, 9)

# Keep the guideline readable in print; body 10.5pt, tables 9pt.
for paragraph in doc.paragraphs:
    if paragraph.style.name in {"Title", "Subtitle", "Heading 1", "Heading 2"}:
        continue
    for run in paragraph.runs:
        if run.font.size and run.font.size.pt < 10:
            run.font.size = Pt(10.5)
        set_east_asia(run, "Hind Siliguri" if any("\u0980" <= c <= "\u09ff" for c in run.text) else "Aptos")

# Explicit production specifications missing from the benchmark.
for paragraph in all_paragraphs(doc):
    if paragraph.text == "Maintain clear space around the outermost visible point of the crest.":
        paragraph.text = "Maintain clear space equal to at least 10% of the crest's overall width on every side, measured from the outermost visible point. Increase space whenever nearby text or edges compete with it."
        set_size(paragraph, 10.5)
    if paragraph.text == "Establish minimum sizes by application: print, signage, website, app icon and embroidery.":
        paragraph.text = "Provisional minimum full-crest sizes pending physical/user testing: 30mm print, 120px digital, 50mm embroidery. Below these, use only an authorized simplified emblem—not a cropped crest. Signage requires site/viewing-distance calculations."
        set_size(paragraph, 10.5)
    if paragraph.text == "Monochrome line icons, subordinate to identity.":
        paragraph.text = "Health Icons outline only; 48-unit source grid; currentColor; English/Bangla labels; CC0 source register; no mixed icon families."
        set_size(paragraph, 9)

# Add accessibility specifics after the palette table.
palette_table = next(t for t in doc.tables if t.cell(0, 0).text == "Role" and t.cell(0, 1).text == "HEX")
new_table = doc.add_table(rows=1, cols=1)
new_table._tbl.getparent().remove(new_table._tbl)
palette_table._tbl.addnext(new_table._tbl)
new_table.cell(0, 0).text = "ACCESSIBILITY / WCAG 2.2 AA target: 4.5:1 normal text; 3:1 large text and meaningful UI components. White on #0076D7 = 4.59:1; white on #005AA8 = 6.94:1. White on #F20D16 = 4.33:1 (not normal-size text); #21B7EA and #FFD21A require dark text or non-text accent use. #64748B on white = 4.76:1. Do not use color alone for clinical or status meaning."
for run in new_table.cell(0, 0).paragraphs[0].runs:
    run.font.size = Pt(10)
    set_east_asia(run, "Aptos")

# Add the supplied South Apollo crest to the identity title page.
if CREST.exists():
    p = doc.paragraphs[0]
    p.alignment = 1
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(CREST), width=Mm(42))
    caption = p.insert_paragraph_before("Supplied institutional crest reference · verify master/vector artwork and rights before production")
    p._p.addnext(caption._p)
    caption.alignment = 1
    for run in caption.runs:
        run.font.size = Pt(9)
        set_east_asia(run, "Aptos")

# Replace the generic legal-name example with the actual approved-wording verification note.
for table in doc.tables:
    for row in table.rows:
        if len(row.cells) > 1 and row.cells[0].text == "Formal name":
            row.cells[1].text = "South Apollo Medical College & Hospital (as supplied)"
            row.cells[2].text = "Verify the current registered college/hospital wording before legal or regulatory use."

# Insert governance/usage addendum before final disclaimer, not on an isolated page.
last = doc.paragraphs[-1]
addendum = last.insert_paragraph_before("PRODUCTION NOTES / v2.2 reviewed copy")
addendum.style = doc.styles["Heading 2"]
notes = [
    "The supplied crest reads SOUTH APOLLO MEDICAL COLLEGE & HOSPITAL. Earlier South Asian references in the concept workspace are superseded for this crest edition; legal-name and rights verification remain required.",
    "The 77-item creative library is a scope checklist, not a claim that all 77 finished assets exist in the supplied package.",
    "Minimum sizes, contrast, icon geometry and repeated table headers are specified in this reviewed copy. Physical print, embroidery, signage, screen-reader and EN/BN user testing remain release gates.",
    "Publication requires a named approver, approval date, source master, rights/consent records, font licenses and a change log. This reviewed document is not itself institutional approval.",
]
for text in notes:
    p = last.insert_paragraph_before(text, style="List Bullet")
    set_size(p, 10)
last.paragraph_format.keep_with_next = False

# Ensure the footer text is updated even if inherited through a separate linked part.
for rel in doc.part.rels.values():
    if rel.reltype.endswith("/footer"):
        for paragraph in rel.target_part.element.xpath(".//w:p"):
            for node in paragraph.xpath(".//w:t"):
                if node.text and "v2.1" in node.text:
                    node.text = node.text.replace("v2.1", "v2.2")

# Apply explicit Bangla fonts in all table content too.
for paragraph in all_paragraphs(doc):
    for run in paragraph.runs:
        if any("\u0980" <= c <= "\u09ff" for c in run.text):
            set_east_asia(run, "Hind Siliguri")

doc.save(OUT)
print("Saved reviewed copy:", OUT)
