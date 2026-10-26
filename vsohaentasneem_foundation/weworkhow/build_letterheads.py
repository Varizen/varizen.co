# -*- coding: utf-8 -*-
"""
STFL Institutional Letterhead Pad Template Generator (Bangla & English)
Conforming strictly to:
- Brand Authority: SohaelTasneem Foundation (STFL)
- Brand Guidelines v3.0.0 & Brand Center Governance
- Canonical Mark: AST-STFL-0001 (stfl-mark-dark.png)
- Official Palette: STFL Deep Ink (#081F2D), Royal Navy (#0B2B43), Sovereign Gold (#C5A059), Stone (#E7E7E2)
- Typography:
    * Bangla Heading: Hind Siliguri (immutable)
    * Bangla Body: Noto Sans Bengali
    * English: Inter
- Platform text ("STFL VI") is intentionally excluded: this is the Foundation's institutional letterhead.
- Single-page A4, clean editorial structure, generous body writing area.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
LOGO_DARK = ROOT.parent / "brandcenter" / "identity" / "monochrome" / "stfl-mark-dark.png"

# Institutional Palette
INK = "081F2D"   # STFL Deep Ink
NAVY = "0B2B43"  # STFL Royal Navy
GOLD = "C5A059"  # STFL Sovereign Gold
STONE = "E7E7E2" # STFL Stone


def set_run_font(run, family, size, color=INK, bold=False, language="en-US"):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), family)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), "bn-BD" if language == "bn-BD" else language)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tcpr.append(margins)


def remove_table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    table._tbl.tblPr.append(borders)


def add_paragraph_border(paragraph, edge, color, size=6, space=2):
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), str(space))
    node.set(qn("w:color"), color)
    borders.append(node)
    ppr.append(borders)


def set_paragraph(paragraph, text, family, size, color=INK, bold=False,
                  align=None, after=0, before=0, language="en-US", line_spacing=1.15):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, family, size, color, bold, language)
    return paragraph


def build_header(section, *, language, org_primary, org_secondary, doc_tag, font_heading):
    header = section.header
    header.is_linked_to_previous = False
    header.distance = Cm(1.2)

    # 3-column table: Logo (left), Institution Titles (middle), Document Tag & Web (right)
    tbl = header.add_table(rows=1, cols=3, width=Cm(17.0))
    tbl.autofit = False
    widths = (Cm(1.8), Cm(10.7), Cm(4.5))
    for c, w in zip(tbl.rows[0].cells, widths):
        c.width = w
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        set_cell_margins(c, top=0, start=0, bottom=60, end=0)
    remove_table_borders(tbl)

    c_logo, c_title, c_meta = tbl.rows[0].cells

    # Left: Official STFL Dark Shield Mark
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_logo.paragraph_format.space_after = Pt(0)
    r_logo = p_logo.add_run()
    r_logo.add_picture(str(LOGO_DARK), width=Cm(1.4))

    # Center: Primary Institution Name + Latin parallel
    p_primary = c_title.paragraphs[0]
    p_primary.paragraph_format.space_before = Pt(0)
    p_primary.paragraph_format.space_after = Pt(1)
    set_paragraph(p_primary, org_primary, font_heading, 17.0, NAVY, bold=True, language=language)

    p_secondary = c_title.add_paragraph()
    p_secondary.paragraph_format.space_before = Pt(0)
    p_secondary.paragraph_format.space_after = Pt(0)
    set_paragraph(p_secondary, org_secondary, "Inter", 9.0, INK, language="en-US")

    # Right: Institutional tag + Canonical Domain
    p_tag = c_meta.paragraphs[0]
    p_tag.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(2)
    set_paragraph(p_tag, doc_tag, font_heading, 9.0, NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, language=language)

    p_url = c_meta.add_paragraph()
    p_url.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_url.paragraph_format.space_before = Pt(0)
    p_url.paragraph_format.space_after = Pt(0)
    set_paragraph(p_url, "stflbd.org", "Inter", 8.5, GOLD, align=WD_ALIGN_PARAGRAPH.RIGHT, language="en-US")

    # Clean double hairline divider: Navy primary + Gold secondary
    rule = header.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(0)
    rule.paragraph_format.line_spacing = 0.3
    add_paragraph_border(rule, "top", NAVY, size=12, space=0)
    add_paragraph_border(rule, "bottom", GOLD, size=6, space=2)


def build_footer(section, *, language, addr_line, comm_line, font_body):
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.distance = Cm(1.2)

    # Hairline divider
    rule = footer.paragraphs[0]
    rule.paragraph_format.space_after = Pt(5)
    rule.paragraph_format.line_spacing = 0.3
    add_paragraph_border(rule, "top", STONE, size=6, space=0)

    # Address
    p_addr = footer.add_paragraph()
    p_addr.paragraph_format.space_before = Pt(2)
    p_addr.paragraph_format.space_after = Pt(1)
    set_paragraph(p_addr, addr_line, font_body, 8.0, INK, language=language)

    # Communication
    p_comm = footer.add_paragraph()
    p_comm.paragraph_format.space_before = Pt(0)
    p_comm.paragraph_format.space_after = Pt(0)
    set_paragraph(p_comm, comm_line, font_body, 8.0, NAVY, language=language)


def build_template(*, filename, language, org_primary, org_secondary, doc_tag,
                   font_heading, font_body, meta_left, meta_right,
                   recipient_label, subject_label, addr_line, comm_line):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.6)
    sec.bottom_margin = Cm(2.4)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    # Base style
    norm = doc.styles["Normal"]
    norm.font.name = font_body
    norm.font.size = Pt(11)
    norm.font.color.rgb = RGBColor.from_string(INK)
    rpr = norm._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_body)

    # Header & Footer
    build_header(sec, language=language, org_primary=org_primary,
                 org_secondary=org_secondary, doc_tag=doc_tag,
                 font_heading=font_heading)
    build_footer(sec, language=language, addr_line=addr_line,
                 comm_line=comm_line, font_body=font_body)

    # Metadata (Reference & Date)
    meta = doc.add_table(rows=1, cols=2)
    meta.autofit = False
    meta.columns[0].width = Cm(8.5)
    meta.columns[1].width = Cm(8.5)
    remove_table_borders(meta)
    for c in meta.rows[0].cells:
        c.width = Cm(8.5)
        set_cell_margins(c, top=0, start=0, bottom=40, end=0)
    set_paragraph(meta.cell(0, 0).paragraphs[0], meta_left, font_body, 9.5, INK, language=language)
    set_paragraph(meta.cell(0, 1).paragraphs[0], meta_right, font_body, 9.5, INK, align=WD_ALIGN_PARAGRAPH.RIGHT, language=language)

    # Recipient
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(14)
    p_rec.paragraph_format.space_after = Pt(10)
    set_paragraph(p_rec, recipient_label, font_body, 10.0, INK, language=language)
    add_paragraph_border(p_rec, "bottom", STONE, size=4, space=3)

    # Subject
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(6)
    p_sub.paragraph_format.space_after = Pt(16)
    set_paragraph(p_sub, subject_label, font_heading, 10.5, NAVY, bold=True, language=language)
    add_paragraph_border(p_sub, "bottom", STONE, size=4, space=3)

    # Properties
    doc.core_properties.title = org_primary
    doc.core_properties.subject = f"{doc_tag} - Official Letterhead Template"
    doc.core_properties.author = "SohaelTasneem Foundation"
    doc.core_properties.keywords = "STFL, letterhead, stationery, template"
    doc.save(ROOT / filename)


if not LOGO_DARK.is_file():
    raise FileNotFoundError(f"Canonical STFL dark mark missing: {LOGO_DARK}")

# 1. Bangla Institutional Letterhead Pad
build_template(
    filename="STFL_Letterhead_Bangla_Template.docx",
    language="bn-BD",
    org_primary="সোহেল তাসনীম ফাউন্ডেশন",
    org_secondary="SohaelTasneem Foundation",
    doc_tag="প্রাতিষ্ঠানিক পত্র",
    font_heading="Hind Siliguri SemiBold",
    font_body="Noto Sans Bengali",
    meta_left="স্মারক নং: ____________________",
    meta_right="তারিখ: ____________________",
    recipient_label="প্রাপক: ______________________________________________________________",
    subject_label="বিষয়: ________________________________________________________________",
    addr_line="এসটিএফএল হাউস, ১৭৩ ই উলান রোড, পশ্চিম রামপুরা, ঢাকা ১২১৭, বাংলাদেশ",
    comm_line="ফোন: +৮৮০১৮৪১০৫০৬৪১১  |  ইমেইল: hi@stflbd.email  |  ওয়েব: stflbd.org",
)

# 2. English Institutional Letterhead Pad
build_template(
    filename="STFL_Letterhead_English_Template.docx",
    language="en-US",
    org_primary="SohaelTasneem Foundation",
    org_secondary="SOHAELTASNEEM FOUNDATION",
    doc_tag="INSTITUTIONAL LETTER",
    font_heading="Inter SemiBold",
    font_body="Inter",
    meta_left="Reference: ____________________",
    meta_right="Date: ____________________",
    recipient_label="To: __________________________________________________________________",
    subject_label="Subject: ______________________________________________________________",
    addr_line="STFL House, 173 E Ullan Road, W Rampura, Dhaka 1217, Bangladesh",
    comm_line="Phone: +88018410506411  |  Email: hi@stflbd.email  |  Web: stflbd.org",
)

print("Generated clean, fully brand-compliant STFL letterhead templates.")
