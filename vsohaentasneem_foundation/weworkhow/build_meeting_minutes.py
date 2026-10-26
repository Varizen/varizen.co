# -*- coding: utf-8 -*-
"""
STFL Meeting Minutes Template generator.
Brand policy: STFL VI Brand Guidelines v3.0.0 + Typography System v2.0.0
- Institutional core palette (Deep Ink / Royal Navy / Snow / Stone / Sovereign Gold / Living Tangerine)
- EN headings: Inter | BN headings: Hind Siliguri | BN body: Noto Sans Bengali
- English first, Bangla directly below
Fonts are embedded downstream via Word COM (embed_meeting_minutes.ps1).
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from PIL import Image

ROOT = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(ROOT, "assets")
MARK_WHITE = os.path.join(ROOT, "..", "main", "web", "public", "stfl-mark-white.png")
MARK_DARK = os.path.join(ROOT, "..", "main", "web", "public", "stfl-mark-dark.png")
OUT_DOCX = os.path.join(ROOT, "STFL_Meeting_Minutes_Template.docx")

os.makedirs(ASSETS, exist_ok=True)

# ---- Brand palette (STFL VI Institutional Core) ----
INK = "081F2D"        # STFL Deep Ink
NAVY = "0B2B43"       # STFL Royal Navy
SNOW = "FAFAF8"       # STFL Snow
STONE = "E7E7E2"      # STFL Stone
GOLD = "C5A059"       # STFL Sovereign Gold
TAN = "FF6B35"        # STFL Living Tangerine
WHITE = "FFFFFF"

# ---- Typography (STFL VI Typography System v2.0.0) ----
EN = "Inter"
EN_MED = "Inter Medium"
EN_SEMI = "Inter SemiBold"
EN_BOLD = "Inter Bold"
BN = "Noto Sans Bengali"
BN_MED = "Noto Sans Bengali Medium"
BN_SEMI = "Noto Sans Bengali SemiBold"
BN_BOLD = "Noto Sans Bengali Bold"
HS = "Hind Siliguri"
HS_SEMI = "Hind Siliguri SemiBold"

# OOXML child-element canonical orders (Word is strict about sequence)
_RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
              "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
              "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz",
              "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
              "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
              "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
              "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
              "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
              "suppressOverlap", "jc", "textDirection", "textAlignment",
              "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
              "pPrChange"]
_TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
                "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
                "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption",
                "tblDescription", "tblPrChange"]
_TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
               "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
               "ins", "del", "trPrChange"]
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
               "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
               "headers", "cellIns", "cellDel", "cellMerge", "tcPrChange"]


def _insert_ordered(parent, el, order):
    tag = el.tag.split("}")[1]
    idx = order.index(tag)
    for child in parent:
        ctag = child.tag.split("}")[1]
        if ctag in order and order.index(ctag) > idx:
            child.addprevious(el)
            return
    parent.append(el)


def _get_or_make(parent, tag, order):
    existing = parent.find(qn("w:" + tag))
    if existing is not None:
        return existing
    el = OxmlElement("w:" + tag)
    _insert_ordered(parent, el, order)
    return el


# ---------------------------------------------------------------- helpers
def style_run(r, en=EN, bn=BN, size=10, bold=False, color=INK, spacing=None):
    r.font.name = en
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:cs"), bn)
    rFonts.set(qn("w:eastAsia"), bn)
    szCs = _get_or_make(rPr, "szCs", _RPR_ORDER)
    szCs.set(qn("w:val"), str(int(size * 2)))
    if bold:
        _get_or_make(rPr, "bCs", _RPR_ORDER)
    if spacing:
        sp = _get_or_make(rPr, "spacing", _RPR_ORDER)
        sp.set(qn("w:val"), str(spacing))


def _fmt(p, before=0, after=0, line=None, align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if keep_next:
        pf.keep_with_next = True
    return pf


def para(container, text="", en=EN, bn=BN, size=10, bold=False, color=INK,
         align=None, before=0, after=0, line=None, spacing=None,
         keep_next=False, rule=None):
    p = container.add_paragraph()
    _fmt(p, before, after, line, align, keep_next)
    if rule:
        par_border(p, rule[0], rule[1], rule[2], rule[3])
    if text:
        r = p.add_run(text)
        style_run(r, en, bn, size, bold, color, spacing)
    return p


def _cell_fmt(p, before=0, after=0, align=None, line=None, rule=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        pf.alignment = align
    if line:
        pf.line_spacing = line
    if rule:
        par_border(p, rule[0], rule[1], rule[2], rule[3])
    return pf


def cell_para(cell, text="", **kw):
    p = cell.paragraphs[0]
    _cell_fmt(p, kw.pop("before", 0), kw.pop("after", 0), kw.pop("align", None),
              kw.pop("line", None), kw.pop("rule", None))
    if text:
        r = p.add_run(text)
        style_run(r, kw.get("en", EN), kw.get("bn", BN), kw.get("size", 10),
                  kw.get("bold", False), kw.get("color", INK), kw.get("spacing"))
    return p


def cell_add(cell, text="", **kw):
    p = cell.add_paragraph()
    _cell_fmt(p, kw.pop("before", 0), kw.pop("after", 0), kw.pop("align", None),
              kw.pop("line", None), kw.pop("rule", None))
    if text:
        r = p.add_run(text)
        style_run(r, kw.get("en", EN), kw.get("bn", BN), kw.get("size", 10),
                  kw.get("bold", False), kw.get("color", INK), kw.get("spacing"))
    return p


def par_border(p, edge, color, sz, space=1):
    pPr = p._p.get_or_add_pPr()
    pBdr = _get_or_make(pPr, "pBdr", _PPR_ORDER)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = _get_or_make(tcPr, "shd", _TCPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)


def table_borders(table, color=STONE, sz=4, nil=False, bottom_only=False):
    tblPr = table._tbl.tblPr
    borders = _get_or_make(tblPr, "tblBorders", _TBLPR_ORDER)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement("w:" + edge)
        if nil or (bottom_only and edge != "bottom"):
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        borders.append(el)


def cell_margins(table, top=60, left=120, bottom=60, right=120):
    tblPr = table._tbl.tblPr
    m = _get_or_make(tblPr, "tblCellMar", _TBLPR_ORDER)
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        m.append(el)


def fix_layout(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = _get_or_make(tblPr, "tblLayout", _TBLPR_ORDER)
    layout.set(qn("w:type"), "fixed")
    ind = _get_or_make(tblPr, "tblInd", _TBLPR_ORDER)
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
            gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


def row_height(row, cm, rule="atLeast"):
    trPr = row._tr.get_or_add_trPr()
    h = _get_or_make(trPr, "trHeight", _TRPR_ORDER)
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), rule)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    _get_or_make(trPr, "cantSplit", _TRPR_ORDER)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = _get_or_make(trPr, "tblHeader", _TRPR_ORDER)
    th.set(qn("w:val"), "true")


def add_field(p, instr, en=EN, size=7, color=INK):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rF = OxmlElement("w:rFonts")
    rF.set(qn("w:ascii"), en)
    rF.set(qn("w:hAnsi"), en)
    rF.set(qn("w:cs"), BN)
    rPr.append(rF)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rPr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size * 2))
    rPr.append(sz)
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(size * 2))
    rPr.append(szCs)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


ANCHOR_TMPL = (
    '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" '
    'relativeHeight="1" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
    '<wp:simplePos x="0" y="0"/>'
    '<wp:positionH relativeFrom="page"><wp:align>center</wp:align></wp:positionH>'
    '<wp:positionV relativeFrom="page"><wp:align>center</wp:align></wp:positionV>'
    '<wp:extent cx="{cx}" cy="{cy}"/>'
    '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
    '<wp:wrapNone/>'
    '<wp:docPr id="{did}" name="STFL Seal Watermark {did}"/>'
    '<wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
    '<pic:pic>'
    '<pic:nvPicPr><pic:cNvPr id="{did}" name="STFL Seal Watermark"/><pic:cNvPicPr/></pic:nvPicPr>'
    '<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
    '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
    '</pic:pic>'
    '</a:graphicData></a:graphic>'
    '</wp:anchor></w:drawing>'
)


def add_watermark(header_par, image_path, width_cm, docpr_id):
    # Inline picture; Word COM converts it to a centered behind-text
    # floating watermark (native Word behavior) in embed_meeting_minutes.ps1.
    run = header_par.add_run()
    run.add_picture(image_path, width=Cm(width_cm))


def make_watermark_png():
    img = Image.open(MARK_DARK).convert("RGBA")
    alpha = img.getchannel("A").point(lambda v: int(v * 0.07))
    img.putalpha(alpha)
    out = os.path.join(ASSETS, "stfl-seal-watermark.png")
    img.save(out)
    return out


def section_header(doc, num_en, num_bn, title_en, title_bn, icon=""):
    p1 = para(doc, "", en=EN_BOLD, bn=HS_SEMI, size=10.5,
              bold=True, color=NAVY, before=12, after=1, keep_next=True)
    if icon:
        ri = p1.add_run(icon + "  ")
        style_run(ri, "Segoe UI Emoji", "Segoe UI Emoji", 10.5, True, NAVY)
    r = p1.add_run(f"{num_en}. {title_en}")
    style_run(r, EN_BOLD, HS_SEMI, 10.5, True, NAVY)
    para(doc, f"{num_bn}. {title_bn}", en=HS_SEMI, bn=HS_SEMI, size=10.5,
         bold=False, color=NAVY, after=5, keep_next=True,
         rule=("bottom", GOLD, 8, 3))


def label_cell(cell, text_en, text_bn, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_para(cell, text_en, en=EN_MED, size=8, bold=False, color=NAVY, align=align, after=0)
    cell_add(cell, text_bn, en=BN, bn=BN, size=8, color=INK, align=align)


def header_cell(cell, text_en, text_bn, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell_para(cell, text_en, en=EN_SEMI, bn=BN_SEMI, size=8.5, bold=False,
              color=WHITE, align=align, after=0)
    cell_add(cell, text_bn, en=BN_SEMI, bn=BN_SEMI, size=8.5, color=WHITE, align=align)


def build_footer(footer):
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    pf = fp.paragraph_format
    pf.space_before = Pt(2)
    pf.tab_stops.add_tab_stop(Cm(8.9), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(17.8), WD_TAB_ALIGNMENT.RIGHT)
    par_border(fp, "top", STONE, 4, 3)
    r = fp.add_run("SohaelTasneem Foundation \u2022 \u09b8\u09cb\u09b9\u09c7\u09b2\u09a4\u09be\u09b8\u09a8\u09c0\u09ae \u09ab\u09be\u0989\u09a8\u09cd\u09a1\u09c7\u09b6\u09a8")
    style_run(r, EN, BN, 7, False, INK)
    r = fp.add_run("\t\u09ae\u09be\u09a8\u09ac\u09bf\u0995 \u09ae\u09b0\u09cd\u09af\u09be\u09a6\u09be \u2022 \u09ac\u09cd\u09af\u09ac\u09b9\u09be\u09b0\u09bf\u0995 \u09aa\u09a6\u0995\u09cd\u09b7\u09c7\u09aa \u2022 \u09aa\u09b0\u09bf\u09ae\u09be\u09aa\u09af\u09cb\u0997\u09cd\u09af \u09aa\u09cd\u09b0\u09ad\u09be\u09ac")
    style_run(r, EN, BN, 6.5, False, INK)
    r = fp.add_run("\tPage ")
    style_run(r, EN, BN, 7, False, INK)
    add_field(fp, "PAGE", EN, 7, INK)
    r = fp.add_run(" of ")
    style_run(r, EN, BN, 7, False, INK)
    add_field(fp, "NUMPAGES", EN, 7, INK)
    r = fp.add_run("  \u2022  STFL-MM-TPL-01 \u2022 v1.0")
    style_run(r, EN, BN, 7, False, INK)


# ---------------------------------------------------------------- build
def build():
    wm_path = make_watermark_png()

    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = EN
    st.font.size = Pt(10)
    st.font.color.rgb = RGBColor.from_string(INK)
    stRpr = st.element.get_or_add_rPr()
    stF = stRpr.get_or_add_rFonts()
    stF.set(qn("w:cs"), BN)
    stF.set(qn("w:eastAsia"), BN)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True

    # ---- headers (watermark on every page; continuation line on page 2+) ----
    fp_header = sec.first_page_header
    fp_header.is_linked_to_previous = False
    add_watermark(fp_header.paragraphs[0], wm_path, 8.6, 9001)

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(2)
    r = hp.add_run("MEETING MINUTES (Continued)  \u2022  ")
    style_run(r, EN_SEMI, BN_SEMI, 8, False, NAVY)
    r2 = hp.add_run("\u09b8\u09ad\u09be\u09b0 \u0995\u09be\u09b0\u09cd\u09af\u09ac\u09bf\u09ac\u09b0\u09a3\u09c0 (\u0985\u09ac\u09cd\u09af\u09be\u09b9\u09a4)")
    style_run(r2, BN_SEMI, BN_SEMI, 8, False, NAVY)
    par_border(hp, "bottom", STONE, 4, 2)
    add_watermark(hp, wm_path, 8.6, 9002)

    # ---- footers (all pages) ----
    build_footer(sec.footer)
    build_footer(sec.first_page_footer)

    # ================= MASTHEAD =================
    mast = doc.add_table(rows=1, cols=3)
    fix_layout(mast, [2.4, 8.4, 7.0])
    table_borders(mast, bottom_only=True)
    borders = mast._tbl.tblPr.find(qn("w:tblBorders"))
    for el in borders:
        if el.tag == qn("w:bottom"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "14")
            el.set(qn("w:color"), GOLD)
    cell_margins(mast, top=140, left=200, bottom=140, right=200)
    row_height(mast.rows[0], 2.1)
    no_split(mast.rows[0])
    for c in mast.rows[0].cells:
        shade_cell(c, NAVY)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    mc = mast.rows[0].cells[0]
    mp = mc.paragraphs[0]
    mp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp.paragraph_format.space_after = Pt(0)
    mp.add_run().add_picture(MARK_WHITE, width=Cm(1.55))

    oc = mast.rows[0].cells[1]
    oc.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = oc.paragraphs[0].add_run("\u09b8\u09cb\u09b9\u09c7\u09b2\u09a4\u09be\u09b8\u09a8\u09c0\u09ae \u09ab\u09be\u0989\u09a8\u09cd\u09a1\u09c7\u09b6\u09a8")
    style_run(r, HS_SEMI, HS_SEMI, 16, False, WHITE)
    cell_add(oc, "\u09ae\u09be\u09a8\u09ac\u09bf\u0995 \u09ae\u09b0\u09cd\u09af\u09be\u09a6\u09be \u2022 \u09ac\u09cd\u09af\u09ac\u09b9\u09be\u09b0\u09bf\u0995 \u09aa\u09a6\u0995\u09cd\u09b7\u09c7\u09aa \u2022 \u09aa\u09b0\u09bf\u09ae\u09be\u09aa\u09af\u09cb\u0997\u09cd\u09af \u09aa\u09cd\u09b0\u09ad\u09be\u09ac",
             en=BN, bn=BN, size=8, color=STONE)

    tc = mast.rows[0].cells[2]
    tc.paragraphs[0].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tc.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = tc.paragraphs[0].add_run("MEETING MINUTES")
    style_run(r, EN_BOLD, HS, 14, False, WHITE, spacing=30)
    cell_add(tc, "\u09b8\u09ad\u09be\u09b0 \u0995\u09be\u09b0\u09cd\u09af\u09ac\u09bf\u09ac\u09b0\u09a3\u09c0",
             en=HS_SEMI, bn=HS_SEMI, size=11.5, color=GOLD, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
    cell_add(tc, "STFL-MM-______-___",
             en=EN_MED, bn=BN, size=8.5, color=TAN, align=WD_ALIGN_PARAGRAPH.RIGHT, after=2)
    cell_add(tc, "OFFICIAL INSTITUTIONAL RECORD \u2022 \u09aa\u09cd\u09b0\u09be\u09a4\u09bf\u09b7\u09cd\u09a0\u09be\u09a8\u09bf\u0995 \u09a8\u09a5\u09bf",
             en=EN, bn=BN, size=6.5, color=STONE, align=WD_ALIGN_PARAGRAPH.RIGHT)

    para(doc, "", size=4, after=4)

    # ================= METADATA =================
    meta = doc.add_table(rows=5, cols=4)
    fix_layout(meta, [4.0, 4.9, 4.0, 4.9])
    table_borders(meta, STONE, 4)
    cell_margins(meta, top=50, left=110, bottom=50, right=110)

    rows = meta.rows
    for row in rows:
        no_split(row)
        row_height(row, 0.95)

    label_cell(rows[0].cells[0], "Meeting Subject", "\u09b8\u09ad\u09be\u09b0 \u09ac\u09bf\u09b7\u09df")
    shade_cell(rows[0].cells[0], SNOW)
    label_cell(rows[0].cells[2], "Meeting No.", "\u09b8\u09ad\u09be \u09a8\u09ae\u09cd\u09ac\u09b0")
    shade_cell(rows[0].cells[2], SNOW)

    label_cell(rows[1].cells[0], "Date", "\u09a4\u09be\u09b0\u09bf\u0996")
    shade_cell(rows[1].cells[0], SNOW)
    label_cell(rows[1].cells[2], "Time", "\u09b8\u09ae\u09df")
    shade_cell(rows[1].cells[2], SNOW)

    label_cell(rows[2].cells[0], "Venue", "\u09b8\u09cd\u09a5\u09be\u09a8")
    shade_cell(rows[2].cells[0], SNOW)
    rows[2].cells[1].merge(rows[2].cells[3])

    label_cell(rows[3].cells[0], "Chairperson", "\u09b8\u09ad\u09be\u09aa\u09a4\u09bf")
    shade_cell(rows[3].cells[0], SNOW)
    label_cell(rows[3].cells[2], "Recording Secretary", "\u09b0\u09c7\u0995\u09b0\u09cd\u09a1\u0995\u09be\u09b0\u09c0 \u09b8\u099a\u09bf\u09ac")
    shade_cell(rows[3].cells[2], SNOW)

    label_cell(rows[4].cells[0], "Meeting Type", "\u09b8\u09ad\u09be\u09b0 \u09a7\u09b0\u09a8")
    shade_cell(rows[4].cells[0], SNOW)
    merged = rows[4].cells[1].merge(rows[4].cells[3])
    p = merged.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("[    ]  Regular \u2014 \u09a8\u09bf\u09df\u09ae\u09bf\u09a4        [    ]  Emergency \u2014 \u099c\u09b0\u09c1\u09b0\u09bf        [    ]  Special \u2014 \u09ac\u09bf\u09b6\u09c7\u09b7")
    style_run(r, EN, BN, 8.5, False, INK)

    # ================= 1. ATTENDANCE =================
    section_header(doc, "1", "\u09e7", "ATTENDANCE", "\u0989\u09aa\u09b8\u09cd\u09a5\u09bf\u09a4\u09bf", icon="\U0001F4CB")

    att = doc.add_table(rows=7, cols=4)
    fix_layout(att, [1.2, 5.6, 6.2, 4.8])
    table_borders(att, STONE, 4)
    cell_margins(att, top=40, left=110, bottom=40, right=110)
    hdr = att.rows[0]
    no_split(hdr)
    repeat_header(hdr)
    row_height(hdr, 0.95)
    heads = [("No.", "\u0995\u09cd\u09b0\u09ae", WD_ALIGN_PARAGRAPH.CENTER),
             ("Name", "\u09a8\u09be\u09ae", None),
             ("Role & Organization", "\u09aa\u09a6\u09ac\u09bf \u0993 \u09aa\u09cd\u09b0\u09a4\u09bf\u09b7\u09cd\u09a0\u09be\u09a8", None),
             ("Signature", "\u09b8\u09cd\u09ac\u09be\u0995\u09cd\u09b7\u09b0", None)]
    for i, (en_t, bn_t, al) in enumerate(heads):
        c = hdr.cells[i]
        shade_cell(c, NAVY)
        header_cell(c, en_t, bn_t, align=al)
    for n in range(1, 7):
        row = att.rows[n]
        no_split(row)
        row_height(row, 0.88)
        c0 = row.cells[0]
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_para(c0, str(n), en=EN, size=9, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        if n % 2 == 0:
            for c in row.cells:
                shade_cell(c, SNOW)

    # ================= 2. AGENDA =================
    section_header(doc, "2", "\u09e8", "AGENDA", "\u0986\u09b2\u09cb\u099a\u09cd\u09af\u09b8\u09c2\u099a\u09c0", icon="\U0001F4DD")

    bn_digits = ["\u09e7", "\u09e8", "\u09e9", "\u09ea", "\u09eb"]
    for i in range(5):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(15)
        pf.space_after = Pt(0)
        par_border(p, "bottom", STONE, 6, 2)
        pf.tab_stops.add_tab_stop(Cm(0.75), WD_TAB_ALIGNMENT.LEFT)
        r = p.add_run(f"{i + 1}.\t{bn_digits[i]}.")
        style_run(r, EN_MED, BN_MED, 9.5, False, NAVY)

    # ---- page break ----
    pb = doc.add_paragraph()
    pb.paragraph_format.space_before = Pt(0)
    pb.paragraph_format.space_after = Pt(0)
    pb.add_run().add_break(WD_BREAK.PAGE)

    # ================= 3. DECISIONS & ACTION ITEMS =================
    section_header(doc, "3", "\u09e9", "DECISIONS & ACTION ITEMS", "\u09b8\u09bf\u09a6\u09cd\u09a7\u09be\u09a8\u09cd\u09a4 \u0993 \u0995\u09be\u09b0\u09cd\u09af \u09ac\u09a3\u09cd\u099f\u09a8", icon="\u26A1")

    dec = doc.add_table(rows=7, cols=5)
    fix_layout(dec, [1.1, 8.1, 3.4, 2.7, 2.5])
    table_borders(dec, STONE, 4)
    cell_margins(dec, top=40, left=110, bottom=40, right=110)
    hdr = dec.rows[0]
    no_split(hdr)
    repeat_header(hdr)
    row_height(hdr, 0.95)
    heads = [("No.", "\u0995\u09cd\u09b0\u09ae", WD_ALIGN_PARAGRAPH.CENTER),
             ("Decision & Action Item", "\u09b8\u09bf\u09a6\u09cd\u09a7\u09be\u09a8\u09cd\u09a4 \u0993 \u0995\u09be\u09b0\u09cd\u09af\u0995\u09cd\u09b0\u09ae", None),
             ("Owner", "\u09a6\u09be\u09af\u09bc\u09bf\u09a4\u09cd\u09ac\u09c7", None),
             ("Deadline", "\u09b8\u09ae\u09af\u09bc\u09b8\u09c0\u09ae\u09be", None),
             ("Status", "\u0985\u09ac\u09b8\u09cd\u09a5\u09be", None)]
    for i, (en_t, bn_t, al) in enumerate(heads):
        c = hdr.cells[i]
        shade_cell(c, NAVY)
        header_cell(c, en_t, bn_t, align=al)
    for n in range(1, 7):
        row = dec.rows[n]
        no_split(row)
        row_height(row, 0.88)
        c0 = row.cells[0]
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_para(c0, str(n), en=EN, size=9, color=INK, align=WD_ALIGN_PARAGRAPH.CENTER)
        if n % 2 == 0:
            for c in row.cells:
                shade_cell(c, SNOW)

    # ================= 4. KEY DISCUSSION POINTS =================
    section_header(doc, "4", "\u09ea", "KEY DISCUSSION POINTS", "\u0986\u09b2\u09cb\u099a\u09a8\u09be\u09b0 \u09ae\u09c2\u09b2 \u09aa\u09cd\u09b0\u09a4\u09bf\u09aa\u09be\u09a6\u09cd\u09af", icon="\U0001F4AC")

    for _ in range(6):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(15)
        pf.space_after = Pt(0)
        par_border(p, "bottom", STONE, 6, 2)
        r = p.add_run(" ")
        style_run(r, EN, BN, 9.5)

    # ================= 5. APPROVAL & SIGNATURES =================
    section_header(doc, "5", "\u09eb", "APPROVAL & SIGNATURES", "\u0985\u09a8\u09c1\u09ae\u09cb\u09a6\u09a8 \u0993 \u09b8\u09cd\u09ac\u09be\u0995\u09cd\u09b7\u09b0", icon="\u270D")

    para(doc, "These minutes are approved as a true and accurate record of the discussions and decisions taken in the meeting.",
         en=EN, bn=BN, size=9, color=INK, after=1, line=1.5)
    para(doc, "\u09b8\u09ad\u09be\u09df \u0997\u09c3\u09b9\u09c0\u09a4 \u0986\u09b2\u09cb\u099a\u09a8\u09be \u0993 \u09b8\u09bf\u09a6\u09cd\u09a7\u09be\u09a8\u09cd\u09a4\u09c7\u09b0 \u09b8\u09a0\u09bf\u0995 \u0993 \u09a8\u09bf\u09b0\u09cd\u09ad\u09c1\u09b2 \u09a8\u09a5\u09bf \u09b9\u09bf\u09b8\u09c7\u09ac\u09c7 \u098f\u0987 \u0995\u09be\u09b0\u09cd\u09af\u09ac\u09bf\u09ac\u09b0\u09a3\u09c0 \u0985\u09a8\u09c1\u09ae\u09cb\u09a6\u09a8 \u0995\u09b0\u09be \u09b9\u09b2\u09cb\u0964",
         en=BN, bn=BN, size=9.5, color=INK, after=8, line=1.7)

    nm = doc.add_table(rows=1, cols=1)
    fix_layout(nm, [17.8])
    table_borders(nm, STONE, 4)
    cell_margins(nm, top=70, left=140, bottom=70, right=140)
    c = nm.rows[0].cells[0]
    shade_cell(c, SNOW)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\U0001F4C5  ")
    style_run(r, "Segoe UI Emoji", "Segoe UI Emoji", 8.5, False, NAVY)
    r = p.add_run("Next Meeting / \u09aa\u09b0\u09ac\u09b0\u09cd\u09a4\u09c0 \u09b8\u09ad\u09be:    ")
    style_run(r, EN_SEMI, BN_SEMI, 8.5, False, NAVY)
    r = p.add_run("Date / \u09a4\u09be\u09b0\u09bf\u0996: ______________     Time / \u09b8\u09ae\u09df: ____________     Venue / \u09b8\u09cd\u09a5\u09be\u09a8: ______________")
    style_run(r, EN, BN, 8.5, False, INK)

    para(doc, "", size=6, after=2)

    sig = doc.add_table(rows=6, cols=2)
    fix_layout(sig, [8.9, 8.9])
    table_borders(sig, nil=True)
    cell_margins(sig, top=30, left=110, bottom=30, right=300)
    r0 = sig.rows[0]
    no_split(r0)
    for col, (en_t, bn_t) in enumerate([
        ("Chairperson\u2019s Signature", "\u09b8\u09ad\u09be\u09aa\u09a4\u09bf\u09b0 \u09b8\u09cd\u09ac\u09be\u0995\u09cd\u09b7\u09b0"),
        ("Recording Secretary\u2019s Signature", "\u09b0\u09c7\u0995\u09b0\u09cd\u09a1\u0995\u09be\u09b0\u09c0 \u09b8\u099a\u09bf\u09ac\u09c7\u09b0 \u09b8\u09cd\u09ac\u09be\u0995\u09cd\u09b7\u09b0"),
    ]):
        c = r0.cells[col]
        cell_para(c, en_t, en=EN_SEMI, bn=BN_SEMI, size=9, color=NAVY, after=0)
        cell_add(c, bn_t, en=HS_SEMI, bn=HS_SEMI, size=9.5, color=NAVY)
    row_height(sig.rows[2], 1.35)
    for col in range(2):
        c = sig.rows[3].cells[col]
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        par_border(p, "bottom", INK, 6, 1)
        r = p.add_run(" ")
        style_run(r, EN, BN, 9.5)
    for col in range(2):
        c = sig.rows[4].cells[col]
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("Name / \u09a8\u09be\u09ae: ______________________________")
        style_run(r, EN, BN, 8.5, False, INK)
    for col in range(2):
        c = sig.rows[5].cells[col]
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("Date / \u09a4\u09be\u09b0\u09bf\u0996: ______________________________")
        style_run(r, EN, BN, 8.5, False, INK)

    cp = doc.core_properties
    cp.title = "Meeting Minutes Template \u2014 \u09b8\u09ad\u09be\u09b0 \u0995\u09be\u09b0\u09cd\u09af\u09ac\u09bf\u09ac\u09b0\u09a3\u09c0"
    cp.author = "SohaelTasneem Foundation (STFL)"
    cp.subject = "Official Institutional Meeting Minutes Template (EN/BN)"
    cp.keywords = "STFL, meeting minutes, \u0995\u09be\u09b0\u09cd\u09af\u09ac\u09bf\u09ac\u09b0\u09a3\u09c0, official template, brand compliant"
    cp.category = "Official Document Template"
    cp.comments = "STFL-MM-TPL-01 v1.0. Brand-compliant bilingual template per STFL VI Brand Guidelines v3.0.0 and Typography System v2.0.0. Fonts embedded."

    doc.save(OUT_DOCX)
    print("saved:", OUT_DOCX)


if __name__ == "__main__":
    build()
